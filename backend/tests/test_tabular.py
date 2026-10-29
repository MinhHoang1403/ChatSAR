"""Kiểm tra parse/normalize với expected độc lập, không lấy output làm gold."""

import io
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import Workbook

from app.errors import AppError
from app.tabular import ColumnReview, profile, read_table, typed_value


def test_csv_bom_delimiter_quotes_and_leading_zero() -> None:
    parsed = read_table('\ufeffcode;label;amount\n0007;"a;b";0.10\n0012;x;0.20\n'.encode(), ".csv")
    assert parsed.headers == ["code", "label", "amount"]
    assert parsed.rows[0] == ["0007", "a;b", "0.10"]
    columns = profile(parsed, uuid4())
    assert columns[0]["data_type"] == "text"
    assert columns[2]["data_type"] == "numeric"


def test_null_empty_zero_are_explicit() -> None:
    base = {"id": "c_" + "a" * 32, "name": "amount", "data_type": "numeric"}
    review = ColumnReview(**base, null_values=[""])
    assert typed_value("", review, "UTC") is None
    assert typed_value("0.00", review, "UTC") == Decimal("0.00")
    text = ColumnReview(id=base["id"], name="code", data_type="text")
    assert typed_value("", text, "UTC") == ""
    assert typed_value("0007", text, "UTC") == "0007"
    assert typed_value(None, text, "UTC") is None


def test_numeric_is_decimal_and_invalid_date_rejected() -> None:
    number = ColumnReview(id="c_" + "b" * 32, name="amount", data_type="numeric")
    assert typed_value("0.10", number, "UTC") + typed_value("0.20", number, "UTC") == Decimal(
        "0.30"
    )
    day = ColumnReview(id="c_" + "c" * 32, name="day", data_type="date")
    with pytest.raises(ValueError):
        typed_value("2026-02-30", day, "UTC")


def test_xlsx_multi_sheet_keeps_numeric_source_and_text() -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Empty"
    first.append(["unused"])
    first.append([1])
    sheet = workbook.create_sheet("Data")
    sheet.append(["code", "amount"])
    sheet.append(["0007", 0.1])
    sheet.append(["0012", None])
    stream = io.BytesIO()
    workbook.save(stream)
    data = stream.getvalue()
    assert read_table(data, ".xlsx").sheets == ["Empty", "Data"]
    parsed = read_table(data, ".xlsx", "Data")
    assert parsed.rows == [["0007", Decimal("0.1")], ["0012", None]]


def test_formula_and_corrupt_excel_rejected() -> None:
    workbook = Workbook()
    workbook.active.append(["amount"])
    workbook.active.append(["=1+1"])
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(AppError):
        read_table(stream.getvalue(), ".xlsx")
    with pytest.raises(AppError):
        read_table(b"corrupt", ".xlsx")


def test_wrong_encoding_and_inconsistent_csv_rejected() -> None:
    with pytest.raises(AppError):
        read_table(b"name\n\xff\n", ".csv")
    with pytest.raises(AppError):
        read_table(b"a,b\n1,2,3\n", ".csv")

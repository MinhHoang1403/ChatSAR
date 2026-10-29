"""Parse bảng local, giữ giá trị nguồn và đề xuất kiểu để người dùng xác nhận."""

from __future__ import annotations

import csv
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Literal
from uuid import UUID, uuid5
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from defusedxml import ElementTree as SafeXML
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from app.errors import AppError

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50000
MAX_COLUMNS = 100
MAX_CELLS = 250000
KINDS = Literal["text", "integer", "numeric", "boolean", "date", "timestamp"]
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


@dataclass
class ParsedTable:
    headers: list[str]
    rows: list[list[Any]]
    sheets: list[str]
    sheet_name: str


class ColumnReview(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True, hide_input_in_errors=True)
    id: str = Field(pattern=r"^c_[a-f0-9]{32}$")
    name: str = Field(min_length=1, max_length=120)
    data_type: KINDS
    nullable: bool = True
    primary_key: bool = False
    null_values: list[str] = Field(default_factory=list, max_length=10)
    date_format: Literal["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"] = "%Y-%m-%d"
    decimal_separator: Literal[".", ","] = "."
    unit: str | None = Field(default=None, max_length=80)
    description: str = Field(default="", max_length=500)
    aliases: list[str] = Field(default_factory=list, max_length=12)
    role: Literal["identifier", "status", "time", "amount", "quantity", "name", "other"] = "other"
    value_mappings: list[ValueMapping] = Field(default_factory=list, max_length=50)
    aggregation_policy: Literal["unspecified", "additive", "non_additive"] = "unspecified"


class ValueMapping(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True, hide_input_in_errors=True)
    phrase: str = Field(min_length=1, max_length=120)
    value: str = Field(max_length=500)


ColumnReview.model_rebuild()


def invalid(message: str) -> AppError:
    return AppError(422, "INVALID_TABULAR_FILE", message)


def read_table(data: bytes, suffix: str, sheet_name: str | None = None) -> ParsedTable:
    if not data or len(data) > MAX_BYTES:
        raise AppError(413, "UPLOAD_LIMIT", "File trống hoặc vượt giới hạn 20 MiB.")
    if suffix == ".csv":
        try:
            encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            text = data.decode(encoding, errors="strict")
            if "\x00" in text or data.startswith(b"PK"):
                raise ValueError("Binary CSV")
            try:
                delimiter = csv.Sniffer().sniff(text[:16384], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
            reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            headers = next(reader)
            if not 1 <= len(headers) <= MAX_COLUMNS:
                raise invalid("Bảng phải có từ 1 đến 100 cột.")
            rows: list[list[Any]] = []
            for row in reader:
                if (
                    len(rows) >= MAX_ROWS
                    or len(row) != len(headers)
                    or (len(rows) + 1) * len(headers) > MAX_CELLS
                ):
                    raise invalid("Số dòng vượt giới hạn hoặc số ô không khớp với tiêu đề.")
                rows.append(row)
        except (UnicodeError, csv.Error, StopIteration, ValueError):
            raise invalid("CSV phải là UTF-8 hoặc UTF-16 có BOM, đúng định dạng bảng.") from None
        parsed = ParsedTable(headers, rows, [], "")
    elif suffix == ".xlsx":
        parsed = read_excel(data, sheet_name)
        if parsed.sheets and not parsed.headers:
            return parsed
    else:
        raise invalid("Chỉ hỗ trợ CSV và XLSX trong giai đoạn này.")
    if not parsed.headers or len(parsed.headers) > MAX_COLUMNS:
        raise invalid("Bảng phải có từ 1 đến 100 cột.")
    if any(not h.strip() or len(h) > 120 for h in parsed.headers):
        raise invalid("Tiêu đề cột trống hoặc quá dài. Vui lòng sửa file nguồn.")
    if len(parsed.rows) * len(parsed.headers) > MAX_CELLS:
        raise invalid("Bảng vượt giới hạn 250.000 ô; vui lòng chia nhỏ file.")
    if not parsed.rows:
        raise invalid("Bảng chưa có dòng dữ liệu.")
    return parsed


def read_excel(data: bytes, sheet_name: str | None) -> ParsedTable:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if (
                len(infos) > 2000
                or sum(i.file_size for i in infos) > 100 * 1024 * 1024
                or any(i.file_size > max(i.compress_size, 1) * 100 for i in infos)
            ):
                raise invalid("XLSX vượt giới hạn giải nén an toàn.")
            workbook_xml = SafeXML.fromstring(archive.read("xl/workbook.xml"))
            sheet_nodes = workbook_xml.findall("s:sheets/s:sheet", NS)
            names = [node.attrib["name"] for node in sheet_nodes]
            if not names or len(names) > 20:
                raise invalid("Workbook trống hoặc có quá nhiều sheet.")
            if sheet_name is None and len(names) > 1:
                return ParsedTable([], [], names, "")
            chosen = sheet_name or names[0]
            if chosen not in names:
                raise invalid("Không tìm thấy sheet đã chọn.")
            relation_id = sheet_nodes[names.index(chosen)].attrib[
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            ]
            relationships = SafeXML.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            target = next(
                node.attrib["Target"] for node in relationships if node.attrib["Id"] == relation_id
            )
            path = (
                target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            )
            if not path.startswith("xl/worksheets/"):
                raise invalid("Sheet path không hợp lệ.")
            numeric: dict[str, Decimal] = {}
            for _, node in SafeXML.iterparse(io.BytesIO(archive.read(path)), events=("end",)):
                if node.tag == "{" + NS["s"] + "}c":
                    if node.attrib.get("t", "n") == "n":
                        value = node.find("s:v", NS)
                        if value is not None and value.text is not None:
                            numeric[node.attrib["r"]] = Decimal(value.text)
                            if len(numeric) > MAX_CELLS + MAX_COLUMNS:
                                raise invalid("Sheet vượt giới hạn số ô.")
                    node.clear()
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            sheet = workbook[chosen]
            sheet.reset_dimensions()
            rows: list[list[Any]] = []
            cell_count = 0
            for cells in sheet.iter_rows():
                if len(rows) > MAX_ROWS or len(cells) > MAX_COLUMNS:
                    raise invalid("Sheet vượt giới hạn dòng hoặc cột.")
                row = []
                for cell in cells:
                    value = cell.value
                    if cell.data_type == "f":
                        raise invalid("Chuyển công thức Excel thành giá trị trước khi tải lên.")
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        if cell.coordinate not in numeric:
                            raise invalid("Không xác minh được giá trị số gốc của Excel.")
                        if re.fullmatch(r"0{2,}", cell.number_format or ""):
                            raise invalid("Mã có zero đầu cần lưu dưới dạng TEXT trong file Excel.")
                        value = numeric[cell.coordinate]
                    row.append(value)
                rows.append(row)
                cell_count += len(row)
                if cell_count > MAX_CELLS + MAX_COLUMNS:
                    raise invalid("Sheet vượt giới hạn số ô.")
        finally:
            workbook.close()
        headers = [str(v) if v is not None else "" for v in rows[0]]
        width = len(headers)
        values = [r + [None] * (width - len(r)) for r in rows[1:]]
        if any(len(r) != width for r in values):
            raise invalid("Dòng Excel có số ô không khớp tiêu đề.")
        return ParsedTable(headers, values, names, chosen)
    except AppError:
        raise
    except Exception:
        # Parser bên ngoài có nhiều loại lỗi; không đưa nội dung workbook ra logs.
        raise invalid("File XLSX hỏng hoặc không đúng định dạng hỗ trợ.") from None


def wire(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if type(value) is int and abs(value) > 9007199254740991:
        return str(value)
    return value


def suggested_type(values: list[Any]) -> KINDS:
    nonempty = [v for v in values if v is not None and v != ""]
    if not nonempty:
        return "text"
    if all(isinstance(v, bool) for v in nonempty):
        return "boolean"
    if all(isinstance(v, (datetime, date)) for v in nonempty):
        return (
            "timestamp"
            if any(isinstance(v, datetime) and v.time().isoformat() != "00:00:00" for v in nonempty)
            else "date"
        )
    strings = [str(v) for v in nonempty]
    if any(re.fullmatch(r"[+-]?0\d+", v) for v in strings):
        return "text"
    if all(re.fullmatch(r"[+-]?\d+", v) for v in strings):
        return "integer"
    try:
        if all(Decimal(v).is_finite() for v in strings):
            return "numeric"
    except InvalidOperation:
        pass
    try:
        if all(re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) and date.fromisoformat(v) for v in strings):
            return "date"
    except ValueError:
        pass
    return "text"


def profile(parsed: ParsedTable, revision_id: UUID) -> list[dict[str, Any]]:
    result = []
    for index, name in enumerate(parsed.headers):
        values = [row[index] for row in parsed.rows]
        nonnull = [wire(v) for v in values if v is not None]
        distinct = {json_key(v) for v in nonnull}
        kind = suggested_type(values)
        result.append(
            {
                "id": "c_" + uuid5(revision_id, str(index)).hex,
                "name": name,
                "data_type": kind,
                "nullable": any(v is None or v == "" for v in values),
                "null_count": sum(v is None for v in values),
                "empty_count": sum(v == "" for v in values),
                "duplicate_count": len(nonnull) - len(distinct),
                "candidate_key": len(nonnull) == len(values) == len(distinct)
                and all(v != "" for v in values),
                "primary_key": False,
                "null_values": [],
                "unit": None,
                "date_format": "%Y-%m-%d",
                "decimal_separator": ".",
            }
        )
    return result


def json_key(value: Any) -> tuple[str, str]:
    return type(value).__name__, str(value)


def typed_value(value: Any, review: ColumnReview, timezone: str) -> Any:
    if value is None or (isinstance(value, str) and value in review.null_values):
        if not review.nullable:
            raise ValueError("NULL not allowed")
        return None
    text = str(value)
    if review.data_type == "text":
        return text
    if review.data_type == "integer":
        if not re.fullmatch(r"[+-]?\d+", text) or not -(2**63) <= int(text) < 2**63:
            raise ValueError("Invalid integer")
        return int(text)
    if review.data_type == "numeric":
        if review.decimal_separator == ",":
            if "." in text:
                raise ValueError("Ambiguous number separators")
            text = text.replace(",", ".")
        number = Decimal(text)
        if not number.is_finite():
            raise ValueError("Invalid number")
        if len(number.as_tuple().digits) > 38 or abs(int(number.as_tuple().exponent)) > 38:
            raise ValueError("Invalid numeric precision")
        return number
    if review.data_type == "boolean":
        if isinstance(value, bool):
            return value
        if text.lower() not in ("true", "false", "1", "0"):
            raise ValueError("Invalid boolean")
        return text.lower() in ("true", "1")
    if review.data_type == "date":
        if isinstance(value, datetime):
            if value.time().isoformat() != "00:00:00":
                raise ValueError("Datetime would lose time")
            return value.date()
        if isinstance(value, date):
            return value
        return datetime.strptime(text, review.date_format).date()
    parsed = value if isinstance(value, datetime) else datetime.fromisoformat(text)
    try:
        zone = ZoneInfo(timezone)
    except ZoneInfoNotFoundError:
        raise ValueError("Invalid timezone") from None
    return parsed.replace(tzinfo=zone) if parsed.tzinfo is None else parsed
